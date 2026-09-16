from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from reportlab.lib import colors #Colores predefinidos para usar en el PDF.
from reportlab.lib.enums import TA_CENTER #Constantes para alinear texto en el PDF (por ejemplo, centrar texto).
from reportlab.lib.pagesizes import letter #Tamaño de página predefinido para el PDF (carta).
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet #Crear estilos de párrafo y obtener estilos de muestra para el PDF.
from reportlab.lib.units import cm #Unidades de medida para el PDF (centímetros).
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle #Elementos para crear el PDF: párrafos, documentos, espacios y tablas.
from reportlab.platypus import Image 



ruta_archivo=Path(__file__).parent / "Registro_Biblioteca.xlsx"

def crear_archivo_excel():
    if not ruta_archivo.exists():
        Nuevo_Archivo_Excel=Workbook()
        hoja_estudiante=Nuevo_Archivo_Excel.active
        hoja_estudiante.title="Estudiantes"
        hoja_libros=Nuevo_Archivo_Excel.create_sheet("Libros")
        hoja_prestamos=Nuevo_Archivo_Excel.create_sheet("Préstamos")
        
        Encabezados = {
            "Estudiantes": ["ID","Nombre","Apellido","Correo","Telefono", "Grado"],
            "Libros": ["ID","Título","Autor","Año","Editorial"],
            "Préstamos": ["ID","ID Estudiante","ID Libro","Fecha Préstamo","Fecha Devolución","Fecha Entrega","Dias Mora","Valor Mora"]
        }
        
        for nombre_hoja, lista_encabezados in Encabezados.items(): #Establecer los encabezados en la primera fila de cada hoja
            hoja_actual = Nuevo_Archivo_Excel[nombre_hoja]
            for columnas, encabezado in enumerate(lista_encabezados, start=1): #Establecer los encabezados en la primera fila de cada hoja
                celda=hoja_actual.cell(row=1, column=columnas, value=encabezado) #Esrablecer el valor de la celda en la primera fila y columna correspondiente al encabezado
                celda.font = Font(bold=True, color="FFFFFF") #Establecer el estilo de fuente en negrita y color blanco para los encabezados
                celda.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid") #Establecer el color de fondo azul para los encabezados
                celda.alignment = Alignment(horizontal="center", vertical="center") ##Alinear el texto de los encabezados al centro horizontal y verticalmente

        Nuevo_Archivo_Excel.save(ruta_archivo)
        print(f"Archivo de Excel creado exitosamente: {ruta_archivo}")
    else:
        Nuevo_Archivo_Excel=load_workbook(ruta_archivo)
        print(f"El archivo ya existe y fue abierto: {ruta_archivo}")
        
    return Nuevo_Archivo_Excel

def Guardar_datos(hoja, datos):
    hoja.append(datos) #Agregar los datos a la hoja de cálculo
    hoja.parent.save(ruta_archivo) #Guardar los cambios en el archivo de Excel
    print("Datos guardados exitosamente.")
    
def buscar_fila(hoja, columna, valor):
    buscado=str(valor).strip().lower() #Convertir el valor a minúsculas y eliminar espacios en blanco al inicio y al final para una búsqueda más flexible.
    for fila in hoja.iter_rows(min_row=2, values_only=True):  # Comenzar desde la fila 2 para omitir los encabezados
        if str(fila[columna]).strip().lower() == buscado: #Comparar el valor de la celda con el valor buscado, ignorando mayúsculas y minúsculas y espacios en blanco
            return fila
    return None
    
def texto_obligatorio(mensaje):
    while True:
        valor = input(mensaje).strip() #Eliminar espacios en blanco al inicio y al final de la entrada del usuario
        if valor: #Si el valor no está vacío, se devuelve el valor ingresado
            return valor
        print("Este campo es obligatorio.")
    
def registrar_estudiante():
    libro_estudiante=load_workbook(ruta_archivo)
    hoja_estudiantes=libro_estudiante["Estudiantes"]
    
    #fila_vacia=hoja_estudiantes.max_row + 1
    
    identificación=texto_obligatorio("Ingrese la identificación del estudiante: ")
    
    fila_existente=buscar_fila(hoja_estudiantes, 0, identificación)
    if fila_existente: #buscar_fila(hoja_estudiantes, 0, identificación):
        print("El estudiante con esa identificación ya existe.")
        return
    
        
    nombre=texto_obligatorio("Ingrese el nombre del estudiante: ")
    apellido=texto_obligatorio("Ingrese el apellido del estudiante: ")
    correo=texto_obligatorio("Ingrese el correo del estudiante: ")
    telefono=texto_obligatorio("Ingrese el teléfono del estudiante: ")
    grado=texto_obligatorio("Ingrese el grado del estudiante: ")
                
    datos_estudiante=[identificación, nombre, apellido, correo, telefono, grado]
        
    """
    for columna, valor in enumerate(datos_estudiante, start=1):
        hoja_estudiantes.cell(row=fila_vacia, column=columna, value=valor)

    libro_estudiante.save(ruta_archivo)
    print("Datos guardados exitosamente.")
    """
    Guardar_datos(hoja_estudiantes, datos_estudiante)
    
def registrar_libro():
    libro_estudiante=load_workbook(ruta_archivo)
    hoja_libros=libro_estudiante["Libros"]
    
    codigo_libro=texto_obligatorio("Ingrese el código del libro: ")     
    
    if buscar_fila(hoja_libros, 0, codigo_libro):
        print("El libro con ese código ya existe.")
        return
       
    titulo=texto_obligatorio("Ingrese el título del libro: ")
    autor=texto_obligatorio("Ingrese el autor del libro: ")
    año=texto_obligatorio("Ingrese el año de publicación del libro: ")
    editorial=texto_obligatorio("Ingrese la editorial del libro: ")
    
    datos_libro=[codigo_libro, titulo, autor, año, editorial]
    
    Guardar_datos(hoja_libros, datos_libro)
    
def generar_id_prestamo(hoja):
    fecha_hoy = date.today().strftime("%d%m%y")

    ultimo_contador = 0

    for fila in hoja.iter_rows(min_row=2, values_only=True):
        valor = str(fila[0]).strip()

        if not valor:
            continue

        # Si la fila tiene el formato: 01 + ddmmyy  -> 01090926
        if len(valor) >= 8 and valor.endswith(fecha_hoy):
            contador = int(valor[:2])
            if contador > ultimo_contador:
                ultimo_contador = contador

    nuevo_contador = ultimo_contador + 1
    return f"{nuevo_contador:02d}{fecha_hoy}"  # Formato: 01 + ddmmyy
    
def registrar_prestamo():
    libro_estudiante=load_workbook(ruta_archivo)
    hoja_prestamos=libro_estudiante["Préstamos"]
    hoja_estudiantes=libro_estudiante["Estudiantes"]
    hoja_libros=libro_estudiante["Libros"]
    
    fecha_prestamo=date.today().strftime("%Y-%m-%d")
    ultimo_contador=0
    
    id_prestamo=generar_id_prestamo(hoja_prestamos)
    
    id_estudiante=texto_obligatorio("Ingrese la identificación del estudiante: ")
    if not buscar_fila(hoja_estudiantes, 0, id_estudiante):
        print("El estudiante no está registrado.")
        return
    
    id_libro=texto_obligatorio("Ingrese el código del libro: ")
    if not buscar_fila(hoja_libros, 0, id_libro):
        print("El libro no está registrado.")
        return
    
    fecha_prestamo=date.today().strftime("%Y-%m-%d")
    fecha_devolucion=texto_obligatorio("Ingrese la fecha de devolución (YYYY-MM-DD): ")
    fecha_entrega=texto_obligatorio("Ingrese la fecha de entrega (YYYY-MM-DD): ")
    
    if fecha_entrega > fecha_devolucion:
        print("La fecha de entrega supera la fecha de devolución")
        dias_mora=(date.fromisoformat(fecha_entrega) - date.fromisoformat(fecha_devolucion)).days
        valor_mora=dias_mora*2000
    else:
        dias_mora=0
        valor_mora=0

    datos_prestamo=[id_prestamo, id_estudiante, id_libro, fecha_prestamo, fecha_devolucion, fecha_entrega, dias_mora, valor_mora]
    
    Guardar_datos(hoja_prestamos, datos_prestamo)
    
def generar_factura(id_prestamo):
    libro=load_workbook(ruta_archivo)
    hoja_prestamos=libro["Préstamos"]
    hoja_estudiantes=libro["Estudiantes"]
    hoja_libros=libro["Libros"]
    
    fila_prestamo=buscar_fila(hoja_prestamos, 0, id_prestamo)
    if not fila_prestamo:
        print("El préstamo con ese ID no existe.")
        return
    
    id_estudiante=str(fila_prestamo[1]).strip() #Permitir que el ID del estudiante sea tratado como una cadena de texto y eliminar espacios en blanco al inicio y al final.
    id_libro=str(fila_prestamo[2]).strip() #Permitir que el ID del libro sea tratado como una cadena de texto y eliminar espacios en blanco al inicio y al final.
    fecha_prestamo=fila_prestamo[3]
    fecha_devolucion=fila_prestamo[4]
    dias_mora=fila_prestamo[6] if len(fila_prestamo) > 6 and fila_prestamo[6] is not None else 0  # Verificar si la fila tiene al menos 7 elementos antes de acceder al índice 6
    valor_mora=fila_prestamo[7] if len(fila_prestamo) > 7 and fila_prestamo[7] is not None else 0  # Verificar si la fila tiene al menos 8 elementos antes de acceder al índice 7
    
    fila_estudiante=buscar_fila(hoja_estudiantes, 0, id_estudiante)
    nombre_estudiante=fila_estudiante[1] if fila_estudiante else "No registrado"
    
    fila_libro=buscar_fila(hoja_libros, 0, id_libro)
    titulo_libro=fila_libro[1] if fila_libro else "No registrado"
    
    fecha_actual=date.today().strftime("%Y-%m-%d")
    
    carpeta_facturas=Path(__file__).parent / "Facturas"
    carpeta_facturas.mkdir(exist_ok=True)
    
    archivo_pdf=carpeta_facturas / f"Factura_{id_prestamo}.pdf"
    
    documento=SimpleDocTemplate(
        str(archivo_pdf), 
        pagesize=letter, rightMargin=2*cm, 
        leftMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    
    estilos=getSampleStyleSheet()
    estilo_titulo=estilos["Title"]
    estilo_titulo.alignment=TA_CENTER
    estilo_texto=estilos["BodyText"]
    
    #logo
    logo_path=Path(__file__).parent / "logo.png"
    if logo_path.exists():
        logo=Image(str(logo_path), width=4*cm, height=4*cm)
    else:
        logo=Paragraph("Logo no disponible", estilo_texto)
        
    contenido=[]
    contenido.append(logo)
    contenido.append(Spacer(1, 12))
    contenido.append(Paragraph("Factura de Préstamo", estilo_titulo))
    contenido.append(Paragraph(f"Fecha: {fecha_actual}", estilo_texto))
    contenido.append(Spacer(1, 12))
    
    datos_tabla=[
        ["Campo", "Valor"],
        ["ID Préstamo",str(id_prestamo)],
        ["ID Estudiante", str(id_estudiante)],
        ["Nombre Estudiante", str(nombre_estudiante)],
        ["ID Libro", str(id_libro)],
        ["Título Libro", str(titulo_libro)],
        ["Fecha Préstamo", str(fecha_prestamo)],
        ["Fecha Devolución", str(fecha_devolucion)],
        ["Días Mora", str(dias_mora)],
        ["Valor Mora", f"${valor_mora:.2f}"]
    ]
    
    tabla=Table(datos_tabla, colWidths=[6*cm, 10*cm])
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (1, 0), colors.white),
        ('ALIGN', (0, 0), (1, 0), 'CENTER'),
        ('FONTNAME', (0, 0), (1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (1, 0), 14),
        ('BOTTOMPADDING', (0, 0), (1, 0), 12),
        ('BACKGROUND', (0, 1), (1, -1), colors.beige),
        ('GRID', (0, 0), (1, -1), 1, colors.black)
    ]))
    contenido.append(tabla)
    
    documento.build(contenido)
    print(f"Factura generada exitosamente: {archivo_pdf}")

def menu_principal():
    crear_archivo_excel()
    while True:
        print("\n--- Menú Principal ---")
        print("1. Registrar estudiante\n2. Registrar libro\n3. Registrar préstamo\n4. Generar Factura\n5. Salir")
        opcion = input("Seleccione una opción: ").strip()
        if opcion == "1":
            registrar_estudiante()
        elif opcion == "2":
            registrar_libro()
        elif opcion == "3":
            registrar_prestamo()
        elif opcion == "4":
            id_prestamo = texto_obligatorio("Ingrese el ID del préstamo: ")
            generar_factura(id_prestamo)
        elif opcion == "5":
            print("Saliendo del programa...")
            break
        else:
            print("Opción no válida. Intente de nuevo.")

menu_principal()
